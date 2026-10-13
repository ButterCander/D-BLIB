import openpyxl, pandas as pd, numpy as np, sys
from scipy.stats import pearsonr
sys.stdout.reconfigure(encoding='utf-8')
def m4():
    wb=openpyxl.load_workbook('si/MOESM4.xlsx',read_only=True)
    out={}
    for ws in wb.worksheets:
        rows=list(ws.iter_rows(values_only=True))
        recs=[]
        for r in rows[2:]:
            try: recs.append(('X%d'%int(r[0]),r[1],r[2],r[3]))
            except: pass
        out[ws.title]=pd.DataFrame(recs,columns=['task','desc','score','p']).set_index('task')
    return out
def m7():
    wb=openpyxl.load_workbook('si/MOESM7.xlsx',read_only=True,data_only=True)
    out={}
    for ws in wb.worksheets:
        if ws.title=='Summary': continue
        rows=list(ws.iter_rows(values_only=True))
        h=rows[0]
        recs=[(r[0],r[2],r[3]) for r in rows[1:] if r[0]]
        out[ws.title]=pd.DataFrame(recs,columns=['task','tide','tidee']).set_index('task')
    return out
a=m4(); b=m7()
for k in b:
    j=a[k].join(b[k],how='inner')
    print(k,len(j),'r(MOESM4 score, MOESM7 tide)=',round(pearsonr(j.score,j.tide)[0],4), 'max abs diff',float(np.abs(j.score-j.tide).max()))
# summary sheet of MOESM7
wb=openpyxl.load_workbook('si/MOESM7.xlsx',read_only=True,data_only=False)
ws=wb['Summary']
for r in list(ws.iter_rows(values_only=True))[:12]: print(r[:9])
print('----')
for k in a:
    for t in ['X76','X113','X138','X140','X141']:
        r4=a[k].loc[t] if t in a[k].index else None
        r7=b[k].loc[t] if t in b[k].index else None
        print(k,t,'M4 score %.3f p %.4f'%(r4.score,r4.p) if r4 is not None else None,'| M7 tide',None if r7 is None else round(float(r7.tide),3), '| desc', None if r4 is None else r4.desc)

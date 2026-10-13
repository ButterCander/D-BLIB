import openpyxl, pandas as pd, numpy as np, sys
from scipy.stats import false_discovery_control
sys.stdout.reconfigure(encoding='utf-8')
wb=openpyxl.load_workbook('si/MOESM4.xlsx',read_only=True)
auth={}
for ws in wb.worksheets:
    rows=list(ws.iter_rows(values_only=True)); recs=[]
    for r in rows[2:]:
        try: recs.append(('X%d'%int(r[0]),r[2],r[3]))
        except: pass
    auth[ws.title]=pd.DataFrame(recs,columns=['task','score','p']).set_index('task')
print('cond  p<.025  BH2p<.05  BH2p<.10 | BH(p tail)<.05 | n p==0 | min p')
for k,d in auth.items():
    p2=np.minimum(2*d.p.values,1.0)
    q=false_discovery_control(p2,method='bh'); qt=false_discovery_control(d.p.values,method='bh')
    print(f'{k:11s} {int((d.p<0.025).sum()):3d}  {int((q<0.05).sum()):3d}  {int((q<0.10).sum()):3d} | {int((qt<0.05).sum()):3d} | {int((d.p==0).sum())} | {d.p.min():.4f}')
# synergy calls BH
syn={}
wb5=openpyxl.load_workbook('si/MOESM5.xlsx',read_only=True,data_only=True)
for nm,cond in [('PI3Ki-MEKi','PI3Ki-MEKi'),('PI3Ki-TAKi','PI3Ki-TAKi')]:
    rows=[r for r in wb5[nm].iter_rows(values_only=True) if any(c is not None for c in r)][1:]
    ids=[r[0] for r in rows]
    d=auth[cond]; p2=np.minimum(2*d.p.values,1.0); q=pd.Series(false_discovery_control(p2,method='bh'),index=d.index)
    qq=q.reindex(ids)
    print(nm,'calls',len(ids),'q<.05',int((qq<0.05).sum()),'q<.10',int((qq<0.10).sum()))
# strong-synergy tasks: single-drug p values and combination q
for nm,single in [('PI3Ki-MEKi','MEKi'),('PI3Ki-TAKi','TAKi')]:
    rows=[r for r in wb5[nm].iter_rows(values_only=True) if any(c is not None for c in r)][1:]
    for r in rows:
        if r[2]==0 and r[3]==0:
            t=r[0]; d=auth[nm]; 
            p2=np.minimum(2*d.p.values,1.0); q=pd.Series(false_discovery_control(p2,method='bh'),index=d.index)
            print(nm,t,r[1][:40],'AB score %.2f p %.4f q %.3f | PI3Ki p %.4f | %s p %.4f'%(r[4],d.p[t],q[t],auth['PI3Ki'].p[t],single,auth[single].p[t]))

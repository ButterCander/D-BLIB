import openpyxl, sys
base = r"D:\Cander\School\3-2\system_bio\mtea\si"
for f in ["MOESM2.xlsx","MOESM3.xlsx","MOESM4.xlsx","MOESM5.xlsx","MOESM7.xlsx"]:
    wb = openpyxl.load_workbook(base+"\\"+f, read_only=True, data_only=True)
    print("==", f)
    for ws in wb.worksheets:
        print("  sheet", repr(ws.title), ws.max_row, ws.max_column)
        for i,row in enumerate(ws.iter_rows(values_only=True)):
            if i<3: print("    ", row[:12])
            else: break

"""Fill the Canva-style template (stage1.pptx: 16 slides already ordered) with the Benedicto 2025 talk."""
import math
import os
import sys
from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "..", "mtea", "figs", "slides")
ORANGE, BLACK, GRAY, MUTED = RGBColor(0xFA, 0x64, 0x3F), RGBColor(0, 0, 0), RGBColor(0xF1, 0xF1, 0xF1), RGBColor(0x6F, 0x6F, 0x6F)
REG, BOLD = "Nanum Gothic", "Nanum Gothic Bold"

prs = Presentation(os.path.join(HERE, "stage1.pptx"))
S = list(prs.slides)


# ---------- helpers ----------
def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)


def find(slide, name):
    for sh in walk(slide.shapes):
        if sh.name == name:
            return sh
    raise KeyError(name)


def remove(sh):
    sh._element.getparent().remove(sh._element)


def _spacing(ppr, line_pct=None, after_pt=None):
    for tag in ("a:lnSpc", "a:spcBef", "a:spcAft"):
        for e in ppr.findall(qn(tag)):
            ppr.remove(e)
    pos = 0
    if line_pct:
        ln = etree.Element(qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPct")).set("val", str(int(line_pct * 1000)))
        ppr.insert(pos, ln)
        pos += 1
    if after_pt is not None:
        bef = etree.Element(qn("a:spcBef"))
        etree.SubElement(bef, qn("a:spcPts")).set("val", "0")
        ppr.insert(pos, bef)
        aft = etree.Element(qn("a:spcAft"))
        etree.SubElement(aft, qn("a:spcPts")).set("val", str(int(after_pt * 100)))
        ppr.insert(pos + 1, aft)


def set_paras(sh, texts, size=None, base_idx=None, bold=None, line_pct=None, after_pt=None):
    """Replace the text of a template shape, keeping each paragraph's run formatting."""
    body = sh.text_frame._txBody
    old = body.findall(qn("a:p"))
    withruns = [p for p in old if p.findall(qn("a:r"))]
    bases = [deepcopy(p) for p in withruns]
    for p in old:
        body.remove(p)
    for i, t in enumerate(texts):
        k = base_idx[i] if base_idx else min(i, len(bases) - 1)
        p = deepcopy(bases[min(k, len(bases) - 1)])
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        for br in p.findall(qn("a:br")):
            p.remove(br)
        r = runs[0]
        r.find(qn("a:t")).text = t
        rpr = r.find(qn("a:rPr"))
        for h in rpr.findall(qn("a:hlinkClick")):
            rpr.remove(h)
        if rpr.get("u"):
            rpr.set("u", "none")
        if size:
            rpr.set("sz", str(int(round(size * 100))))
        if bold is not None:
            rpr.set("b", "true" if bold else "false")
        if line_pct or after_pt is not None:
            ppr = p.find(qn("a:pPr"))
            if ppr is None:
                ppr = etree.Element(qn("a:pPr"))
                p.insert(0, ppr)
            _spacing(ppr, line_pct, after_pt)
        body.append(p)


def add_text(slide, x, y, w, h, paras, size, bold=False, color=BLACK, align=None, bullets=False, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = t
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.name = BOLD if bold else REG
        f.color.rgb = color
        if bullets:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(int(Inches(0.32))))
            ppr.set("indent", str(-int(Inches(0.32))))
            bf = etree.SubElement(ppr, qn("a:buFont"))
            bf.set("typeface", "Arial")
            bc = etree.SubElement(ppr, qn("a:buChar"))
            bc.set("char", "•")
    return tb


def add_line(slide, x1, y, x2, color=ORANGE, width=1.5):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


def add_rect(slide, x, y, w, h, fill=GRAY):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def add_title(slide, num, title, x=1.44, size=70, w=16.0):
    add_text(slide, x, 1.14, 1.4, 0.5, [num], 25, bold=True, color=ORANGE)
    add_text(slide, x, 1.62, w, 1.4, [title], size, bold=True)


def pill(slide, group_name, text_name, label):
    set_paras(find(slide, text_name), [label])


def lines_needed(text, size, width_in, bold=False):
    cpl = width_in / ((0.62 if bold else 0.56) * size / 72.0)
    return max(1, math.ceil(len(text) / cpl))


def side_block(slide, x, y, w, heading, items, hsize=26, bsize=21):
    n = lines_needed(heading, hsize, w, True)
    hh = n * hsize * 1.3 / 72
    add_text(slide, x, y, w, hh, [heading], hsize, bold=True)
    ly = y + hh + 0.15
    add_line(slide, x, ly, x + w)
    add_text(slide, x, ly + 0.25, w, 4.5, items, bsize, bullets=True, space_after=11)


def picture_h(slide, name, x, y, h):
    return slide.shapes.add_picture(os.path.join(FIG, name), Inches(x), Inches(y), height=Inches(h))


def picture_w(slide, name, x, y, w):
    return slide.shapes.add_picture(os.path.join(FIG, name), Inches(x), Inches(y), width=Inches(w))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def fix_figure_slide(slide, num, label):
    """Strip the template's photo layout from a copy of slide 11; keep the pill."""
    for n in ("AutoShape 2", "Picture 3", "Group 8", "Group 12"):
        remove(find(slide, n))
    pill(slide, "Group 4", "TextBox 7", label)


def table_el(slide):
    return find(slide, "Table 3")._element.find(".//" + qn("a:tbl"))


def reshape_table(tbl, widths_emu, nrows, header_h, body_h):
    grid = tbl.find(qn("a:tblGrid"))
    cols = grid.findall(qn("a:gridCol"))
    while len(cols) < len(widths_emu):
        grid.append(deepcopy(cols[-1]))
        cols = grid.findall(qn("a:gridCol"))
    for g, w in zip(cols, widths_emu):
        g.set("w", str(int(w)))
    trs = tbl.findall(qn("a:tr"))
    for tr in trs:
        tcs = tr.findall(qn("a:tc"))
        while len(tcs) < len(widths_emu):
            tr.append(deepcopy(tcs[-1]))
            tcs = tr.findall(qn("a:tc"))
    body = trs[1]
    while len(tbl.findall(qn("a:tr"))) < nrows:
        tbl.append(deepcopy(body))
    for i, tr in enumerate(tbl.findall(qn("a:tr"))):
        tr.set("h", str(int(header_h if i == 0 else body_h)))


def set_cell(tc, texts, size, plain=False, align=None):
    body = tc.find(qn("a:txBody"))
    if plain:
        tcpr = tc.find(qn("a:tcPr"))
        tcpr.set("marT", "45720")
        tcpr.set("marB", "45720")
    ps = body.findall(qn("a:p"))
    base = deepcopy([p for p in ps if p.findall(qn("a:r"))][0])
    for p in ps:
        body.remove(p)
    for t in texts:
        p = deepcopy(base)
        runs = p.findall(qn("a:r"))
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = t
        runs[0].find(qn("a:rPr")).set("sz", str(int(size * 100)))
        ppr = p.find(qn("a:pPr"))
        if plain and ppr is not None:
            _spacing(ppr, 100, None)
            for a in ("marL", "indent", "lvl"):
                if a in ppr.attrib:
                    del ppr.attrib[a]
            for tag in ("a:buFont", "a:buChar"):
                for e in ppr.findall(qn(tag)):
                    ppr.remove(e)
            etree.SubElement(ppr, qn("a:buNone"))
            # buNone must precede defRPr
            d = ppr.find(qn("a:defRPr"))
            if d is not None:
                ppr.remove(d)
                ppr.append(d)
        if align and ppr is not None:
            ppr.set("algn", align)
        body.append(p)
        # a:txBody children order: bodyPr, lstStyle, p...
    return tc


def fill_table(tbl, rows, header_size, body_size, plain_cols=(), align_cols=None):
    trs = tbl.findall(qn("a:tr"))
    for ri, row in enumerate(rows):
        tcs = trs[ri].findall(qn("a:tc"))
        for ci, val in enumerate(row):
            texts = val if isinstance(val, list) else [val]
            if ri == 0:
                set_cell(tcs[ci], texts, header_size, plain=False)
            else:
                al = align_cols[ci] if align_cols else None
                set_cell(tcs[ci], texts, body_size, plain=(ci in plain_cols), align=al)


# ---------- 1 cover ----------
s = S[0]
t = find(s, "TextBox 10")
t.width = Inches(10.4)
set_paras(t, ["Constraint-based modeling of drug-induced metabolic changes in a cancer cell line"], size=44, line_pct=105)
set_paras(find(s, "TextBox 6"), ["Benedicto et al., 2025"])
set_paras(find(s, "TextBox 8"), ["Systems Biology Midterm"])
set_paras(find(s, "TextBox 14"), ["권태민 · 32220304"])
set_paras(find(s, "TextBox 13"), ["Format B · Presentation date 2026-10-14"])
notes(s, "Cue: npj Systems Biology and Applications 11:111 (2025). AGS gastric cancer cells, three kinase inhibitors, two combinations. Public data GSE285616 and open-source MTEApy; I re-ran the analysis myself.")

# ---------- 2 outline ----------
s = S[1]
set_paras(find(s, "TextBox 6"), ["Outline"])
items = {
    "TextBox 9": "Background", "TextBox 11": "Method", "TextBox 14": "Result 1 · Fig. 3", "TextBox 17": "Result 2 · Fig. 4",
    "TextBox 10": "Cross-check · Fig. 5", "TextBox 12": "My re-analysis", "TextBox 15": "Assessment", "TextBox 18": "Relevance",
    "TextBox 8": "Take-home", "TextBox 13": "Questions", "TextBox 16": "Backup · Fig. 1, Fig. 2", "TextBox 19": "Backup · Table 1, terms",
}
for n, txt in items.items():
    box = find(s, n)
    box.width = Inches(4.6)
    set_paras(box, [txt])
set_paras(find(s, "TextBox 5"), ["Paper: Benedicto, Flobak, Ponce-de-León, Valencia", "npj Systems Biology and Applications 11:111 (2025)", "Data: GEO GSE285616 · Code: MTEApy (open source)"], base_idx=[0, 1, 2], size=17)
notes(s, "Optional slide; skip if time is short.")

# ---------- 3 background ----------
s = S[2]
set_paras(find(s, "TextBox 5"), ["Background"])
set_paras(find(s, "TextBox 10"), ["Background"])
set_paras(
    find(s, "TextBox 11"),
    [
        "Cancer cells rewire metabolism under PI3K, MEK and TAK1 signaling",
        "PI3Ki+MEKi and PI3Ki+TAKi inhibit AGS growth synergistically",
        "Gap: expression analysis gives broad labels, not metabolic detail",
        "Question: which metabolic tasks change, and which only in the combination?",
    ],
    size=22, line_pct=100, after_pt=9,
)
remove(find(s, "Group 6"))
add_rect(s, 12.9, 1.12, 6.0, 9.0)
add_text(s, 13.4, 1.7, 5.0, 0.6, ["The system"], 27, bold=True)
add_line(s, 13.4, 2.45, 18.4)
add_text(
    s, 13.4, 2.75, 5.0, 6.0,
    ["AGS gastric adenocarcinoma cells", "PI3K inhibitor PI-103 (PI3Ki)", "MEK inhibitor PD0325901 (MEKi)", "TAK1 inhibitor (5Z)-7-oxozeaenol (TAKi)", "Two combinations: PI3Ki+MEKi, PI3Ki+TAKi", "Control: DMSO; 4 replicates per condition"],
    21, bullets=True, space_after=14,
)
notes(s, "Cue: AGS carries KRAS G12D and PIK3CA mutations, so MAPK and PI3K are both active. Synergy of the pairs was shown for growth in Flobak 2015. GSEA and DEG give broad labels such as amino acid metabolism.")

# ---------- 4 method (3 steps) ----------
s = S[3]
remove(find(s, "Group 11"))
add_title(s, "04", "Method", x=1.44)
pill(s, "Group 15", "TextBox 18", "Method")
set_paras(find(s, "TextBox 4"), ["1  RNA-seq and DESeq2", "6 conditions × 4 replicates", "log2FC against DMSO control"], size=22)
set_paras(find(s, "TextBox 7"), ["2  Metabolic map", "Human-GEM and 187 tasks", "AND = min, OR = absmax", "Task score = mean log2FC"], size=22)
set_paras(find(s, "TextBox 10"), ["3  Test and rule", "10,000 permutations, p < 0.025", "Weak synergy: beyond both single drugs", "Strong synergy: only the combination"], size=22)
notes(s, "Cue: a task is a defined input to output conversion, e.g. aspartate to asparagine. TIDE averages the reactions of a pFBA reference pathway; TIDE-essential averages genes whose knockout makes the task infeasible. Permutation shuffles gene log2FC values.")

# ---------- 5 Fig 3 ----------
s = S[4]
fix_figure_slide(s, "05", "Result 1 · Fig. 3")
add_title(s, "05", "Result 1")
picture_w(s, "fig3.png", 1.12, 3.1, 12.4)
add_text(s, 1.12, 9.1, 12.4, 0.4, ["Source: Fig. 3, Benedicto et al. 2025 (CC BY-NC-ND 4.0). A tasks · B overlap · C systems · D shared tasks"], 14, color=MUTED)
side_block(
    s, 13.9, 3.1, 5.0, "Metabolic tasks are weakened, led by PI3K inhibition",
    ["77, 66 and 55 significant tasks in PI3Ki, PI3Ki+MEKi and PI3Ki+TAKi", "About 88% of significant tasks are weakened, not strengthened", "Asp → Asn synthesis: −5.9 in PI3Ki, weakened in all five conditions"],
)
notes(s, "Cue: A = significant tasks per condition (down in blue). B = Jaccard overlap, PI3Ki and PI3Ki+MEKi 0.75, MEKi isolated. C = amino acids largest. D = tasks significant in all five; Asp to Asn is ASNS, a single-gene task. Task score = mean log2FC.")

# ---------- 6 Fig 4 ----------
s = S[5]
fix_figure_slide(s, "06", "Result 2 · Fig. 4")
add_title(s, "06", "Result 2")
picture_h(s, "fig4_AC.png", 1.12, 3.0, 6.5)
side_block(
    s, 8.7, 3.1, 10.2, "Some tasks are weakened only by the combination", hsize=28, bsize=22, items=
    [
        "49 of 66 PI3Ki+MEKi tasks are classified synergistic (PI3Ki+TAKi: 14 of 55)",
        "Weak synergy: the combination is more extreme than both single drugs",
        "Strong synergy: only the combination is significant, e.g. spermidine synthesis, CoQ synthesis, thioredoxin reductase",
        "Most calls are weak synergy (about 43 of 49)",
    ],
)
add_text(s, 8.7, 9.4, 10.2, 0.4, ["Source: Fig. 4A and 4C, Benedicto et al. 2025 (CC BY-NC-ND 4.0)"], 14, color=MUTED)
notes(s, "Cue: parent and child logic. Weak: combination below the min of both singles. Strong: both singles not significant, combination significant. Strong-synergy scores are small, -0.3 to -0.8, versus -5.9 for Asp to Asn. Combination is half-dose of each drug.")

# ---------- 7 Fig 5 ----------
s = S[6]
fix_figure_slide(s, "07", "Cross-check · Fig. 5")
add_title(s, "07", "Cross-check")
picture_w(s, "fig5.png", 1.12, 3.1, 12.0)
add_text(s, 1.12, 8.55, 12.0, 0.4, ["Source: Fig. 5, Benedicto et al. 2025 (CC BY-NC-ND 4.0). A overlap · B score agreement · C TIDE-essential tasks"], 14, color=MUTED)
side_block(
    s, 13.7, 3.1, 5.2, "A second method: TIDE-essential",
    ["245 vs 69 significant task–condition pairs (TIDE vs TIDE-essential)", "40 shared; Spearman ρ = 0.93 on these shared pairs", "New in TIDE-essential: β-alanine synthesis, down in both combinations"],
)
notes(s, "Cue: counts are task-condition pairs, not distinct tasks. rho = 0.93 is computed on pairs already significant in both methods. Single-gene tasks give identical scores. Complementary, not a validation.")

# ---------- 8 re-analysis ----------
s = S[7]
fix_figure_slide(s, "08", "My re-analysis")
add_title(s, "08", "My re-analysis")
picture_h(s, "repro_scatter_w.png", 1.12, 3.1, 4.7)
picture_h(s, "repro_counts_w.png", 6.75, 3.1, 4.7)
add_text(s, 1.12, 7.95, 5.3, 0.4, ["Scores: six tasks × five conditions"], 15, color=MUTED)
add_text(s, 6.75, 7.95, 6.4, 0.4, ["Counts depend on the permutation null"], 15, color=MUTED)
add_text(s, 1.12, 8.85, 12.0, 0.4, ["Data: GEO GSE285616; code: MTEApy; pydeseq2 and 1,000 permutations (paper: R DESeq2, 10,000)"], 14, color=MUTED)
side_block(
    s, 13.7, 3.1, 5.2, "Scores reproduce; counts depend on method choices",
    ["DEG counts within about 10% of Fig. 1A", "Task scores agree: r = 0.994 (30 values)", "TIDE vs TIDE-essential: ρ 0.91 on shared pairs, 0.61 on all pairs"],
)
notes(s, "Cue: same counts from GEO, pydeseq2 with replicate + condition, apeglm, p < 0.025. Null = shuffle log2FC among model genes only or among all genes; the paper's ordering PI3Ki, PI3Ki+MEKi, PI3Ki+TAKi, TAKi, MEKi appears only with all genes. Trust direction and scores, not exact counts.")

# ---------- 9 assessment (table, 3 columns) ----------
s = S[8]
remove(find(s, "Group 8"))
add_title(s, "09", "Assessment", x=1.44)
pill(s, "Group 4", "TextBox 7", "Assessment")
tbl = table_el(s)
total = 15633756
reshape_table(tbl, [total / 3] * 3, 2, 1103465, 3215007)
fill_table(
    tbl,
    [
        ["Robust", "Sensitive", "Not tested"],
        [
            ["Direction: tasks are weakened, most under PI3K inhibition", "Scores reproduce from public data", "Asp → Asn (ASNS) is the clearest signal"],
            ["Task counts depend on the permutation null", "Weak synergy: margin over the best single drug is often below 0.1 log2FC", "Strong synergy flips near p = 0.025"],
            ["Multiple testing across 187 tasks", "Interaction term: combination vs single drugs", "Experimental validation; one cell line, one dose"],
        ],
    ],
    header_size=26, body_size=20,
)
add_text(s, 1.44, 10.35, 17.1, 0.5, ["Printed totals in Table 1 do not match their own columns (TAKi 25 vs 24 + 6; MEKi TIDE-essential 15 vs 9 + 7); corrected values are used."], 14, color=MUTED)
notes(s, "Cue: robust = what I could reproduce. Sensitive = what moved when I changed the null or looked at margins (median margin 0.06 to 0.10; 52 to 70 percent of weak calls below 0.1). Not tested = acknowledged by the authors or absent.")

# ---------- 10 relevance (T12) ----------
s = S[9]
set_paras(find(s, "TextBox 5"), ["Relevance"])
set_paras(find(s, "TextBox 15"), ["Cancer and", "drug discovery"], size=54)
set_paras(find(s, "TextBox 16"), ["10"])
for n, txt in (
    ("TextBox 21", ["Combination therapy: a metabolic readout for why PI3K + MEK co-inhibition limits growth more than PI3K + TAK1"]),
    ("TextBox 22", ["Candidate vulnerabilities: ODC1 (polyamines), ASNS (asparagine), thioredoxin reductase; hypotheses only"]),
    ("TextBox 23", ["Scalable readout: RNA-seq, a public model and open code can be applied to other lines, drugs and doses"]),
):
    box = find(s, n)
    box.width = Inches(5.7)
    set_paras(box, txt, size=20)
add_line(s, 1.99, 6.2, 8.9)
add_text(s, 1.99, 6.45, 7.0, 1.6, ["Next step I would take: multiple-testing control, an interaction model, dose–response data"], 22)
notes(s, "Cue: AGS has KRAS G12D and PIK3CA mutations, so both arms are on. ODC1 is the target of DFMO; asparagine synthesis links to asparaginase therapy; thioredoxin to CHK1 inhibitor sensitivity (cited in the paper). All hypotheses, from background knowledge.")

# ---------- 11 take-home (T7) ----------
s = S[10]
pill(s, "Group 3", "TextBox 6", "Take-home")
remove(find(s, "TextBox 8"))
from pptx.enum.text import PP_ALIGN
add_text(s, 1.8, 3.35, 16.4, 3.0, ["PI3K inhibition weakens biosynthetic tasks.", "Combination-only changes are candidates, not proofs."], 34, bold=True, align=PP_ALIGN.CENTER, space_after=14)
set_paras(find(s, "TextBox 9"), ["11"])
set_paras(find(s, "TextBox 10"), ["Scores reproduce from public data; counts and synergy labels depend on method choices."], size=24)
notes(s, "Three points: (1) PI3K inhibition weakens most biosynthetic tasks, asparagine synthesis clearest. (2) Combination-only changes such as polyamine and CoQ synthesis are candidates. (3) Scores reproduce; counts and synergy labels depend on method choices.")

# ---------- 12 thanks (T14) ----------
s = S[11]
set_paras(find(s, "TextBox 4"), ["Thank you"])
set_paras(find(s, "TextBox 5"), ["12"])
set_paras(find(s, "TextBox 6"), ["Questions?"])

# ---------- 13 backup Fig 1 ----------
s = S[12]
fix_figure_slide(s, "13", "Backup · Fig. 1")
add_title(s, "13", "Backup")
picture_w(s, "fig1.png", 1.12, 3.1, 12.0)
add_text(s, 1.12, 8.75, 12.0, 0.4, ["Source: Fig. 1, Benedicto et al. 2025 (CC BY-NC-ND 4.0)"], 14, color=MUTED)
side_block(
    s, 13.7, 3.1, 5.2, "Gene-level response (Fig. 1)",
    ["Up-regulated genes outnumber down-regulated in every condition", "PI3Ki+MEKi has the most DEGs", "Combination-only DEGs: about 15% (PI3Ki+TAKi), about 25% (PI3Ki+MEKi)"],
)

# ---------- 14 backup Fig 2 ----------
s = S[13]
fix_figure_slide(s, "14", "Backup · Fig. 2")
add_title(s, "14", "Backup")
picture_h(s, "fig2.png", 1.12, 2.95, 7.0)
add_text(s, 1.12, 10.05, 8.0, 0.4, ["Source: Fig. 2, Benedicto et al. 2025 (CC BY-NC-ND 4.0)"], 14, color=MUTED)
side_block(
    s, 7.9, 3.1, 11.0, "Gene-set response (Fig. 2)",
    [
        "About 54% of significant GO sets are down, 46% up; PI3Ki combinations lean down",
        "Down in all conditions: rRNA processing, tRNA aminoacylation, mitochondrial translation",
        "Up in all conditions: MHC antigen presentation",
        "Labels stay broad: no pathway-level metabolic detail, hence the model-based step",
        "My check: no single mitochondrial ribosomal gene reaches 2-fold, yet the set is strongly down (p about 1e-19)",
    ],
)

# ---------- 15 backup Table 1 ----------
s = S[14]
remove(find(s, "Group 8"))
add_title(s, "15", "Table 1, corrected", x=1.44, size=60)
pill(s, "Group 4", "TextBox 7", "Backup · Table 1")
tbl = table_el(s)
w0 = Inches(4.5)
rest = (total - w0) / 5
reshape_table(tbl, [w0] + [rest] * 5, 6, Inches(0.85), Inches(0.62))
fill_table(
    tbl,
    [
        ["Significant tasks", "TAKi", "MEKi", "PI3Ki", "PI3Ki+MEKi", "PI3Ki+TAKi"],
        ["TIDE only", "24", "10", "70", "52", "49"],
        ["TIDE-essential only", "8", "9", "3", "4", "5"],
        ["Shared", "6", "7", "7", "14", "6"],
        ["Total TIDE", "30 *", "17", "77", "66", "55"],
        ["Total TIDE-essential", "14", "16 *", "10", "18", "11"],
    ],
    header_size=22, body_size=20, plain_cols=(0, 1, 2, 3, 4, 5), align_cols=["l", "ctr", "ctr", "ctr", "ctr", "ctr"],
)
add_text(s, 1.44, 9.95, 17.1, 0.6, ["* The paper prints 25 and 15. Rows add up to 24 + 6 = 30 and 9 + 7 = 16, which matches the text and Fig. 5A (245 and 69 task–condition pairs)."], 15, color=MUTED)

# ---------- 16 backup terms ----------
s = S[15]
remove(find(s, "Group 8"))
add_title(s, "16", "Terms for questions", x=1.44, size=60)
pill(s, "Group 4", "TextBox 7", "Backup · Terms")
tbl = table_el(s)
reshape_table(tbl, [Inches(4.2), total - Inches(4.2)], 7, Inches(0.8), Inches(0.62))
fill_table(
    tbl,
    [
        ["Term", "Meaning"],
        ["GPR rule", "Complex = AND (take min); isozymes = OR (take largest |log2FC|)"],
        ["pFBA", "Flux balance analysis picking the smallest-total-flux pathway per task"],
        ["FVA", "Flux range per reaction; 149 of 187 tasks have alternative pathways"],
        ["Permutation null", "Shuffle gene log2FC 10,000 times; significant if tail fraction < 0.025"],
        ["GI50", "Dose that halves growth; singles at GI50, combinations at half of each"],
        ["Jaccard index", "Set overlap: intersection divided by union (0 to 1)"],
    ],
    header_size=22, body_size=20, plain_cols=(0, 1), align_cols=["l", "l"],
)

# ---------- cleanup: hyperlinks and unused relationships ----------
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
for slide in prs.slides:
    root = slide._element
    for h in list(root.iter(qn("a:hlinkClick"))):
        h.getparent().remove(h)
    used = set()
    for el in root.iter():
        for k, v in el.attrib.items():
            if k.startswith(RID):
                used.add(v)
    for rid, rel in list(slide.part.rels.items()):
        if rid in used:
            continue
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        slide.part.drop_rel(rid)

out = os.path.join(HERE, "..", "Benedicto2025_midterm_template.pptx")
prs.save(out)
print("saved", os.path.abspath(out))
